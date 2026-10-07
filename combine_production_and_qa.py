import os
import re
import json
import datetime
import shutil
from copy import copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BACKUP_SOURCE = r'D:\Output like same\1. PRODUCTION SKILL ASSESSMENT DATA 30.06.2026.backup.xlsx'
OUTPUT_TARGET = r'D:\Output like same\1. PRODUCTION SKILL ASSESSMENT DATA 30.06.2026.xlsx'
DOWNLOADS_TARGET = r'C:\Users\ReubenG\Downloads\1. PRODUCTION SKILL ASSESSMENT DATA 30.06.2026.xlsx'

print('Loading backup workbook:', BACKUP_SOURCE)
wb = openpyxl.load_workbook(BACKUP_SOURCE, data_only=False)

# -------------------------------------------------------------
# 1. LOAD QA EMPLOYEES & SCORES
# -------------------------------------------------------------
with open('data.js', encoding='utf-8') as f:
    txt = f.read()

m = re.search(r'const EMPLOYEES = (\[.*?\]);', txt, re.DOTALL)
employees = json.loads(m.group(1))
print(f'Loaded {len(employees)} QA employees from data.js')

records = {}
if os.path.exists('assessment_records_backup_236.json'):
    with open('assessment_records_backup_236.json', encoding='utf-8') as f:
        records.update(json.load(f))
if os.path.exists('assessment_records.json'):
    with open('assessment_records.json', encoding='utf-8') as f:
        records.update(json.load(f))
print(f'Loaded assessment records: {len(records)}')

ojt_records = {}
if os.path.exists('ojt_evaluations.json'):
    try:
        with open('ojt_evaluations.json', encoding='utf-8') as f:
            ojt_records = json.load(f)
    except Exception:
        ojt_records = {}
print(f'Loaded OJT records: {len(ojt_records)}')

STANDARD_SECTIONS = [
    ('TB QA Lead', 'Tire building QA'),
    ('FF QA Lead', 'Final Finish QA'),
    ('PR QA Lead', 'Preparatory QA'),
    ('ST QA Lead', 'Solid tire QA'),
    ('WH QA Lead', 'Warehouse QA'),
    ('TC QA Lead', 'Tire curing QA'),
    ('RRO & ALT Lead', 'Final Finish RRO & ALT QA'),
    ('FID QA Lead', 'FID inspector QA')
]

def map_standard_section(s):
    if not s: return 'Final Finish QA'
    sl = s.lower().strip()
    if 'ware' in sl: return 'Warehouse QA'
    if 'build' in sl: return 'Tire building QA'
    if 'curing' in sl: return 'Tire curing QA'
    if 'solid' in sl: return 'Solid tire QA'
    if 'prep' in sl: return 'Preparatory QA'
    if 'rro' in sl or 'alt' in sl: return 'Final Finish RRO & ALT QA'
    if 'fid' in sl: return 'FID inspector QA'
    return 'Final Finish QA'

def sort_key(emp):
    sec = map_standard_section(emp.get('section', ''))
    sec_order = [s[1] for s in STANDARD_SECTIONS]
    sec_idx = sec_order.index(sec) if sec in sec_order else 99
    return (sec_idx, emp.get('empNo', ''))

sorted_qa_employees = sorted(employees, key=sort_key)

# -------------------------------------------------------------
# 2. APPEND QA EMPLOYEES TO 'Name list' SHEET
# -------------------------------------------------------------
ws_nl = wb['Name list']
print(f'Current max row in Name list: {ws_nl.max_row} (Rows 3 to 1181 are Production)')

# Verify row 1181 is intact
assert ws_nl.cell(1181, 4).value == 'PRODUCTION', 'Row 1181 is not PRODUCTION!'

# Cache style template from row 1181
style_template = {}
for col in range(1, 27):
    c1181 = ws_nl.cell(1181, col)
    style_template[col] = {
        'font': copy(c1181.font),
        'border': copy(c1181.border),
        'alignment': copy(c1181.alignment),
        'fill': copy(c1181.fill),
        'number_format': c1181.number_format
    }

start_qa_row = 1182
for i, emp in enumerate(sorted_qa_employees):
    row_idx = start_qa_row + i
    emp_no = emp.get('empNo', '')
    name = emp.get('name', '')
    dept = 'QUALITY CONTROL'
    section = map_standard_section(emp.get('section', ''))
    doj_str = emp.get('doj', '2020-01-01')
    try:
        doj_date = datetime.datetime.strptime(doj_str, '%Y-%m-%d')
    except Exception:
        doj_date = datetime.datetime(2020, 1, 1)

    cur_level = emp.get('currentLevel') or 'I'
    if cur_level == 'I': eligible_level = 'L'
    elif cur_level == 'L': eligible_level = 'U'
    elif cur_level == 'U': eligible_level = 'O'
    else: eligible_level = 'O'

    rec = records.get(emp_no, {})
    has_exam = rec.get('isCompleted') or (rec.get('totalMark') is not None and rec.get('totalMark') > 0)
    
    sq = rec.get('submittedQuestions', [])
    cat_correct = {'Safety': 0, 'CI & TPM': 0, 'Process': 0, 'Quality': 0, 'System': 0}
    
    if sq:
        for q in sq:
            cat = (q.get('category') or '').lower()
            if q.get('isCorrect'):
                if 'safety' in cat:
                    cat_correct['Safety'] += 1
                elif 'ci' in cat or 'tpm' in cat:
                    cat_correct['CI & TPM'] += 1
                elif 'system' in cat:
                    cat_correct['System'] += 1
                elif 'process' in cat and 'quality' in cat:
                    cat_correct['Process'] += 0.5
                    cat_correct['Quality'] += 0.5
                elif 'process' in cat:
                    cat_correct['Process'] += 1
                else:
                    cat_correct['Quality'] += 1
        safety_score = int(round(cat_correct['Safety']))
        ci_score = int(round(cat_correct['CI & TPM']))
        proc_score = int(round(cat_correct['Process']))
        qual_score = int(round(cat_correct['Quality']))
        sys_score = int(round(cat_correct['System'])) if cat_correct['System'] > 0 else None
    elif has_exam:
        tm = rec.get('totalMark', 0)
        safety_score = min(10, int(tm * 0.25))
        proc_score = min(10, int(tm * 0.25))
        qual_score = min(10, int(tm * 0.25))
        ci_score = max(0, tm - safety_score - proc_score - qual_score)
        sys_score = None
    else:
        safety_score = None
        proc_score = None
        qual_score = None
        ci_score = None
        sys_score = None

    ojt = ojt_records.get(emp_no, {})
    if ojt.get('scorePct') is not None or ojt.get('totalScore') is not None:
        gemba_s = int(ojt.get('safetyScore') or 15)
        gemba_sop = int(ojt.get('sopScore') or 15)
    elif has_exam and rec.get('status') == 'Passed':
        gemba_s = 15
        gemba_sop = 15
    elif has_exam:
        gemba_s = 12
        gemba_sop = 10
    else:
        gemba_s = None
        gemba_sop = None

    # Values
    ws_nl[f'A{row_idx}'] = f'=ROW()-2'
    ws_nl[f'B{row_idx}'] = emp_no
    ws_nl[f'C{row_idx}'] = name
    ws_nl[f'D{row_idx}'] = dept
    ws_nl[f'E{row_idx}'] = section
    ws_nl[f'F{row_idx}'] = doj_date
    ws_nl[f'G{row_idx}'] = f'=NOW()-F{row_idx}'
    ws_nl[f'H{row_idx}'] = eligible_level
    ws_nl[f'I{row_idx}'] = cur_level
    ws_nl[f'J{row_idx}'] = safety_score
    ws_nl[f'K{row_idx}'] = proc_score
    ws_nl[f'L{row_idx}'] = qual_score
    ws_nl[f'M{row_idx}'] = ci_score
    ws_nl[f'N{row_idx}'] = sys_score
    ws_nl[f'O{row_idx}'] = f'=SUM(J{row_idx}:N{row_idx})'
    ws_nl[f'P{row_idx}'] = f'=O{row_idx}/IF(H{row_idx}="L",20,IF(H{row_idx}="U",30,IF(H{row_idx}="O",40,IF(H{row_idx}="I",20,""))))'
    ws_nl[f'Q{row_idx}'] = f'=IF(AND(H{row_idx}="L", P{row_idx}>=50%), "Pass", IF(AND(H{row_idx}="U", P{row_idx}>=60%), "Pass", IF(AND(H{row_idx}="O", P{row_idx}>=75%), "Pass", "Fail")))'
    ws_nl[f'R{row_idx}'] = None
    ws_nl[f'S{row_idx}'] = gemba_s
    ws_nl[f'T{row_idx}'] = gemba_sop
    ws_nl[f'U{row_idx}'] = f'=SUM(S{row_idx}:T{row_idx})'
    ws_nl[f'V{row_idx}'] = f'=U{row_idx}/IF(H{row_idx}="L",35,IF(H{row_idx}="U",35,IF(H{row_idx}="O",35,IF(H{row_idx}="I",35,""))))'
    ws_nl[f'W{row_idx}'] = f'=IF(AND(H{row_idx}="L", V{row_idx}>=50%), "Pass", IF(AND(H{row_idx}="U", V{row_idx}>=60%), "Pass", IF(AND(H{row_idx}="O", V{row_idx}>=75%), "Pass", "Fail")))'
    ws_nl[f'X{row_idx}'] = f'=IF(AND(Q{row_idx}="Pass", W{row_idx}="Pass"), "Pass", "Fail")'
    ws_nl[f'Y{row_idx}'] = f'=IF($X{row_idx}="Pass",$H{row_idx}, IF(H{row_idx}="L","I", IF(H{row_idx}="U","L", IF(H{row_idx}="O","U", "I"))))'
    ws_nl[f'Z{row_idx}'] = None

    # Apply style from row 1181 template
    for col in range(1, 27):
        c = ws_nl.cell(row_idx, col)
        st = style_template[col]
        c.font = copy(st['font'])
        c.border = copy(st['border'])
        c.alignment = copy(st['alignment'])
        c.fill = copy(st['fill'])
        c.number_format = st['number_format']
    
    # Specific number formats
    ws_nl[f'F{row_idx}'].number_format = 'dd/mmm/yyyy'
    ws_nl[f'G{row_idx}'].number_format = 'yy\\-mm'
    ws_nl[f'P{row_idx}'].number_format = '0%'
    ws_nl[f'V{row_idx}'].number_format = '0%'

final_nl_row = start_qa_row + len(sorted_qa_employees) - 1
print(f'Name list updated: appended QA rows {start_qa_row} to {final_nl_row}. Total employees: {final_nl_row - 2}')

# -------------------------------------------------------------
# 3. UPDATE 'Abstract' SHEET
# -------------------------------------------------------------
ws_abs = wb['Abstract']

# Unmerge lower ranges (e.g. G31:H31)
for m in list(ws_abs.merged_cells.ranges):
    if m.min_row >= 29:
        ws_abs.unmerge_cells(str(m))

# Cache styles from row 28 (section row), row 29 (Grand Total), row 30 (%), row 31 (U&O)
abs_sec_style = {c: {'font': copy(ws_abs.cell(28, c).font), 'border': copy(ws_abs.cell(28, c).border), 'fill': copy(ws_abs.cell(28, c).fill), 'alignment': copy(ws_abs.cell(28, c).alignment), 'num_fmt': ws_abs.cell(28, c).number_format} for c in range(1, 20)}
abs_tot_style = {c: {'font': copy(ws_abs.cell(29, c).font), 'border': copy(ws_abs.cell(29, c).border), 'fill': copy(ws_abs.cell(29, c).fill), 'alignment': copy(ws_abs.cell(29, c).alignment), 'num_fmt': ws_abs.cell(29, c).number_format} for c in range(1, 20)}
abs_pct_style = {c: {'font': copy(ws_abs.cell(30, c).font), 'border': copy(ws_abs.cell(30, c).border), 'fill': copy(ws_abs.cell(30, c).fill), 'alignment': copy(ws_abs.cell(30, c).alignment), 'num_fmt': ws_abs.cell(30, c).number_format} for c in range(1, 20)}
abs_uo_style  = {c: {'font': copy(ws_abs.cell(31, c).font), 'border': copy(ws_abs.cell(31, c).border), 'fill': copy(ws_abs.cell(31, c).fill), 'alignment': copy(ws_abs.cell(31, c).alignment), 'num_fmt': ws_abs.cell(31, c).number_format} for c in range(1, 20)}

# Populate 8 QA sections at rows 29 to 36
for i, (lead, sec_name) in enumerate(STANDARD_SECTIONS):
    r = 29 + i
    ws_abs[f'B{r}'] = lead
    ws_abs[f'C{r}'] = sec_name
    ws_abs[f'D{r}'] = f"=COUNTIFS('Name list'!$E:$E,Abstract!$C{r})"
    ws_abs[f'E{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,Abstract!E$3,'Name list'!$E:$E,Abstract!$C{r})"
    ws_abs[f'F{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,Abstract!F$3,'Name list'!$E:$E,Abstract!$C{r})"
    ws_abs[f'G{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,Abstract!G$3,'Name list'!$E:$E,Abstract!$C{r})"
    ws_abs[f'H{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,Abstract!H$3,'Name list'!$E:$E,Abstract!$C{r})"
    ws_abs[f'I{r}'] = None
    ws_abs[f'J{r}'] = f"=COUNTIFS('Name list'!$E$3:$E$1048576,Abstract!$C{r},'Name list'!$O$3:$O$1048576,\">1\")"
    ws_abs[f'K{r}'] = f"=J{r}/D{r}"
    ws_abs[f'L{r}'] = f"=COUNTIFS('Name list'!$E$3:$E$1048576,Abstract!$C{r},'Name list'!$U$3:$U$1048576,\">1\")"
    ws_abs[f'M{r}'] = f"=L{r}/D{r}"
    ws_abs[f'N{r}'] = f"=AVERAGE(K{r},M{r})"
    ws_abs[f'O{r}'] = None
    ws_abs[f'P{r}'] = f"=D{r}-J{r}"
    ws_abs[f'Q{r}'] = f"=D{r}-L{r}"

    for col in range(2, 18):
        c = ws_abs.cell(r, col)
        st = abs_sec_style[col]
        c.font = copy(st['font'])
        c.border = copy(st['border'])
        c.fill = copy(st['fill'])
        c.alignment = copy(st['alignment'])
        c.number_format = st['num_fmt']
    ws_abs[f'K{r}'].number_format = '0%'
    ws_abs[f'M{r}'].number_format = '0%'
    ws_abs[f'N{r}'].number_format = '0%'

# Row 37: Grand Total
abs_tot_r = 37
ws_abs[f'B{abs_tot_r}'] = None
ws_abs[f'C{abs_tot_r}'] = 'Grand Total'
ws_abs[f'D{abs_tot_r}'] = f'=SUM(D4:D36)'
ws_abs[f'E{abs_tot_r}'] = f'=SUM(E4:E36)'
ws_abs[f'F{abs_tot_r}'] = f'=SUM(F4:F36)'
ws_abs[f'G{abs_tot_r}'] = f'=SUM(G4:G36)'
ws_abs[f'H{abs_tot_r}'] = f'=SUM(H4:H36)'
ws_abs[f'I{abs_tot_r}'] = None
ws_abs[f'J{abs_tot_r}'] = f'=SUM(J4:J36)'
ws_abs[f'K{abs_tot_r}'] = f'=AVERAGE(K4:K36)'
ws_abs[f'L{abs_tot_r}'] = f'=SUM(L4:L36)'
ws_abs[f'M{abs_tot_r}'] = f'=AVERAGE(M4:M36)'
ws_abs[f'N{abs_tot_r}'] = f'=AVERAGE(N4:N36)'
ws_abs[f'O{abs_tot_r}'] = None
ws_abs[f'P{abs_tot_r}'] = f'=SUM(P4:P36)'
ws_abs[f'Q{abs_tot_r}'] = f'=SUM(Q4:Q36)'

for col in range(2, 18):
    c = ws_abs.cell(abs_tot_r, col)
    st = abs_tot_style[col]
    c.font = copy(st['font'])
    c.border = copy(st['border'])
    c.fill = copy(st['fill'])
    c.alignment = copy(st['alignment'])
    c.number_format = st['num_fmt']
ws_abs[f'K{abs_tot_r}'].number_format = '0%'
ws_abs[f'M{abs_tot_r}'].number_format = '0%'
ws_abs[f'N{abs_tot_r}'].number_format = '0%'

# Row 38: %
abs_pct_r = 38
for c in range(1, 18):
    ws_abs.cell(abs_pct_r, c).value = None
ws_abs[f'E{abs_pct_r}'] = f'=E{abs_tot_r}/$D${abs_tot_r}'
ws_abs[f'F{abs_pct_r}'] = f'=F{abs_tot_r}/$D${abs_tot_r}'
ws_abs[f'G{abs_pct_r}'] = f'=G{abs_tot_r}/$D${abs_tot_r}'
ws_abs[f'H{abs_pct_r}'] = f'=H{abs_tot_r}/$D${abs_tot_r}'

for col in range(5, 9):
    c = ws_abs.cell(abs_pct_r, col)
    st = abs_pct_style[col]
    c.font = copy(st['font'])
    c.border = copy(st['border'])
    c.fill = copy(st['fill'])
    c.alignment = copy(st['alignment'])
    c.number_format = '0%'

# Row 39: U & O combined
abs_uo_r = 39
for c in range(1, 18):
    ws_abs.cell(abs_uo_r, c).value = None
ws_abs[f'G{abs_uo_r}'] = f'=H{abs_pct_r}+G{abs_pct_r}'
ws_abs.merge_cells(f'G{abs_uo_r}:H{abs_uo_r}')
for col in [7, 8]:
    c = ws_abs.cell(abs_uo_r, col)
    st = abs_uo_style[col]
    c.font = copy(st['font'])
    c.border = copy(st['border'])
    c.fill = copy(st['fill'])
    c.alignment = copy(st['alignment'])
    c.number_format = '0%'

print('Abstract sheet updated successfully with rows 4-28 (Production), 29-36 (QA), 37 (Total), 38 (%), 39 (U&O).')

# -------------------------------------------------------------
# 4. UPDATE '2025-2026 comp' SHEET
# -------------------------------------------------------------
ws_comp = wb['2025-2026 comp']

# Unmerge lower ranges (e.g. K32:L32)
for m in list(ws_comp.merged_cells.ranges):
    if m.min_row >= 30:
        ws_comp.unmerge_cells(str(m))

# Cache styles from row 29 (section row), row 30 (Grand Total), row 31 (%), row 32 (U&O)
comp_sec_style = {c: {'font': copy(ws_comp.cell(29, c).font), 'border': copy(ws_comp.cell(29, c).border), 'fill': copy(ws_comp.cell(29, c).fill), 'alignment': copy(ws_comp.cell(29, c).alignment), 'num_fmt': ws_comp.cell(29, c).number_format} for c in range(1, 20)}
comp_tot_style = {c: {'font': copy(ws_comp.cell(30, c).font), 'border': copy(ws_comp.cell(30, c).border), 'fill': copy(ws_comp.cell(30, c).fill), 'alignment': copy(ws_comp.cell(30, c).alignment), 'num_fmt': ws_comp.cell(30, c).number_format} for c in range(1, 20)}
comp_pct_style = {c: {'font': copy(ws_comp.cell(31, c).font), 'border': copy(ws_comp.cell(31, c).border), 'fill': copy(ws_comp.cell(31, c).fill), 'alignment': copy(ws_comp.cell(31, c).alignment), 'num_fmt': ws_comp.cell(31, c).number_format} for c in range(1, 20)}
comp_uo_style  = {c: {'font': copy(ws_comp.cell(32, c).font), 'border': copy(ws_comp.cell(32, c).border), 'fill': copy(ws_comp.cell(32, c).fill), 'alignment': copy(ws_comp.cell(32, c).alignment), 'num_fmt': ws_comp.cell(32, c).number_format} for c in range(1, 20)}

# Populate 8 QA sections at rows 30 to 37
for i, (lead, sec_name) in enumerate(STANDARD_SECTIONS):
    r = 30 + i
    ws_comp[f'B{r}'] = lead
    ws_comp[f'C{r}'] = sec_name
    ws_comp[f'D{r}'] = f"=COUNTIFS('Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'E{r}'] = f"=COUNTIFS('Name list'!$I:$I,'2025-2026 comp'!E$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'F{r}'] = f"=COUNTIFS('Name list'!$I:$I,'2025-2026 comp'!F$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'G{r}'] = f"=COUNTIFS('Name list'!$I:$I,'2025-2026 comp'!G$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'H{r}'] = f"=COUNTIFS('Name list'!$I:$I,'2025-2026 comp'!H$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'I{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,'2025-2026 comp'!I$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'J{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,'2025-2026 comp'!J$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'K{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,'2025-2026 comp'!K$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'L{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,'2025-2026 comp'!L$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'M{r}'] = None
    ws_comp[f'N{r}'] = f'=I{r}-E{r}'
    ws_comp[f'O{r}'] = f'=J{r}-F{r}'
    ws_comp[f'P{r}'] = f'=K{r}-G{r}'
    ws_comp[f'Q{r}'] = f'=L{r}-H{r}'

    for col in range(2, 18):
        c = ws_comp.cell(r, col)
        st = comp_sec_style[col]
        c.font = copy(st['font'])
        c.border = copy(st['border'])
        c.fill = copy(st['fill'])
        c.alignment = copy(st['alignment'])
        c.number_format = st['num_fmt']

# Row 38: Grand Total in 2025-2026 comp
comp_tot_r = 38
ws_comp[f'B{comp_tot_r}'] = None
ws_comp[f'C{comp_tot_r}'] = 'Grand Total'
ws_comp[f'D{comp_tot_r}'] = f'=SUM(D4:D37)'
ws_comp[f'E{comp_tot_r}'] = f'=SUM(E4:E37)'
ws_comp[f'F{comp_tot_r}'] = f'=SUM(F4:F37)'
ws_comp[f'G{comp_tot_r}'] = f'=SUM(G4:G37)'
ws_comp[f'H{comp_tot_r}'] = f'=SUM(H4:H37)'
ws_comp[f'I{comp_tot_r}'] = f'=SUM(I4:I37)'
ws_comp[f'J{comp_tot_r}'] = f'=SUM(J4:J37)'
ws_comp[f'K{comp_tot_r}'] = f'=SUM(K4:K37)'
ws_comp[f'L{comp_tot_r}'] = f'=SUM(L4:L37)'

for col in range(2, 18):
    c = ws_comp.cell(comp_tot_r, col)
    st = comp_tot_style[col]
    c.font = copy(st['font'])
    c.border = copy(st['border'])
    c.fill = copy(st['fill'])
    c.alignment = copy(st['alignment'])
    c.number_format = st['num_fmt']

# Row 39: % in 2025-2026 comp
comp_pct_r = 39
for c in range(1, 18):
    ws_comp.cell(comp_pct_r, c).value = None
ws_comp[f'E{comp_pct_r}'] = f'=E{comp_tot_r}/$D${comp_tot_r}'
ws_comp[f'F{comp_pct_r}'] = f'=F{comp_tot_r}/$D${comp_tot_r}'
ws_comp[f'G{comp_pct_r}'] = f'=G{comp_tot_r}/$D${comp_tot_r}'
ws_comp[f'H{comp_pct_r}'] = f'=H{comp_tot_r}/$D${comp_tot_r}'
ws_comp[f'I{comp_pct_r}'] = f'=I{comp_tot_r}/$D${comp_tot_r}'
ws_comp[f'J{comp_pct_r}'] = f'=J{comp_tot_r}/$D${comp_tot_r}'
ws_comp[f'K{comp_pct_r}'] = f'=K{comp_tot_r}/$D${comp_tot_r}'
ws_comp[f'L{comp_pct_r}'] = f'=L{comp_tot_r}/$D${comp_tot_r}'

for col in range(5, 13):
    c = ws_comp.cell(comp_pct_r, col)
    st = comp_pct_style[col]
    c.font = copy(st['font'])
    c.border = copy(st['border'])
    c.fill = copy(st['fill'])
    c.alignment = copy(st['alignment'])
    c.number_format = '0%'

# Row 40: U & O combined in 2025-2026 comp
comp_uo_r = 40
for c in range(1, 18):
    ws_comp.cell(comp_uo_r, c).value = None
ws_comp[f'K{comp_uo_r}'] = f'=L{comp_pct_r}+K{comp_pct_r}'
ws_comp.merge_cells(f'K{comp_uo_r}:L{comp_uo_r}')
for col in [11, 12]:
    c = ws_comp.cell(comp_uo_r, col)
    st = comp_uo_style[col]
    c.font = copy(st['font'])
    c.border = copy(st['border'])
    c.fill = copy(st['fill'])
    c.alignment = copy(st['alignment'])
    c.number_format = '0%'

print('2025-2026 comp sheet updated successfully with rows 4-29 (Production), 30-37 (QA), 38 (Total), 39 (%), 40 (U&O).')

# -------------------------------------------------------------
# 5. SAVE WORKBOOK
# -------------------------------------------------------------
wb.save(OUTPUT_TARGET)
print(f'Successfully saved combined workbook to: {OUTPUT_TARGET}')

try:
    shutil.copy2(OUTPUT_TARGET, DOWNLOADS_TARGET)
    print(f'Mirrored to downloads: {DOWNLOADS_TARGET}')
except Exception as e:
    print('Error copying to downloads:', e)
