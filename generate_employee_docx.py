import os
import sys
import re
import json
import argparse
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
QC_DIR = os.path.join(BASE_DIR, "QC question")
if not os.path.exists(QC_DIR):
    alt_qc = r"C:\Users\ReubenG\Downloads\ILUO MCQ TO Excel sheet\QC question"
    if os.path.exists(alt_qc):
        QC_DIR = alt_qc

DATA_JS = os.path.join(BASE_DIR, "data.js")
LOGO_PATH = os.path.join(BASE_DIR, "yokohama_logo.png")
if not os.path.exists(LOGO_PATH):
    alt_logo = r"C:\Users\ReubenG\Downloads\ILUO MCQ TO Excel sheet\yokohama_logo.png"
    if os.path.exists(alt_logo):
        LOGO_PATH = alt_logo

OJT_JSON = os.path.join(BASE_DIR, "ojt_evaluations.json")
RECORDS_JSON = os.path.join(BASE_DIR, "assessment_records.json")
OUTPUT_DIR = os.path.join(BASE_DIR, "output_docx")

os.makedirs(OUTPUT_DIR, exist_ok=True)

# Brand Colors
COLOR_NAVY = RGBColor(0, 43, 73)
COLOR_YOKO_BLUE = RGBColor(0, 91, 158)
COLOR_GREEN = RGBColor(22, 101, 52)
COLOR_RED = RGBColor(185, 28, 28)
COLOR_MUTED = RGBColor(100, 116, 139)
COLOR_DARK = RGBColor(15, 23, 42)

TEMPLATE_MAP = {
    ('L', 'final finish rro & alt qa'): 'L Level - Final Finish RRO & ALT QA.docx',
    ('L', 'tire building qa'): 'L Level - Tire Building QA.docx',
    ('L', 'tire curing qa'): 'L Level - Tire Curing QA.docx',
    ('L', 'solid tire qa'): 'L Level - Solid Tire QA.docx',
    ('L', 'preparatory qa'): 'L Level - Preparatory QA.docx',
    ('L', 'fid inspector qa'): 'L Level - FID Inspector QA.docx',
    ('L', 'warehouse qa'): 'L Level - Warehouse QA.docx',
    ('L', 'final finish qa'): 'L Level - Final Finish QA.docx',

    ('U', 'final finish rro & alt qa'): 'U Level - Final Finish RRo & ALT QA.docx',
    ('U', 'tire building qa'): 'U Level - Tire Building QA.docx',
    ('U', 'tire curing qa'): 'U Level - Tire Curing QA.docx',
    ('U', 'solid tire qa'): 'U Level - Solid Tire QA.docx',
    ('U', 'preparatory qa'): 'U Level - Preparatory QA.docx',
    ('U', 'fid inspector qa'): 'U Level - FID Inspector QA.docx',
    ('U', 'warehouse qa'): 'U Level - Warehouse QA.docx',
    ('U', 'final finish qa'): 'U Level - Final Finish QA.docx',

    ('O', 'final finish rro & alt qa'): 'O Level - Final Finish RRO & ALT QA.docx',
    ('O', 'tire building qa'): 'O Level - Tire Building QA.docx',
    ('O', 'tire curing qa'): 'O Level - Tire Curing QA.docx',
    ('O', 'solid tire qa'): 'O Level - Solid Tire QA.docx',
    ('O', 'preparatory qa'): 'O Level - Preparatory QA.docx',
    ('O', 'fid inspector qa'): 'O Level - FID Inspector QA.docx',
    ('O', 'warehouse qa'): 'O Level - Warehouse QA.docx',
    ('O', 'final finish qa'): 'O Level - Final Finish QA.docx',
}

OJT_CHECKPOINTS = {
    'tire_building': {
        'formatNo': 'Format No: ATC/T/FOR/HR/87D',
        'title': 'ON THE JOB TRAINING EVALUATION - TBM QA',
        'checkpoints': [
            (1, 'Safety Awareness (BBs, PPE, Tire handling, Electrical safety,)'),
            (2, 'Machine cleanliness and material handling'),
            (3, 'Verification for BPR parameter & centering.'),
            (4, 'Verification of Drum parameter with filling of the drum change memo and FTC sheet.'),
            (5, 'Bottom / Back stitcher tool gap and play verification.'),
            (6, 'Verification on – NSNL /MES/SKU sticker and recipe /Guide light / Pressure gauge/ material Guider centering/ TCU temperature /Pokayoke'),
            (7, 'Material measurement and knowledge on measuring tool'),
            (8, 'Material direction (Ply/breaker/belt) or material positioning for uncommon size.'),
            (9, 'CC/ GT defect checking and NC material handling'),
            (10, 'TEI ( Knowledge of QCC, 5S, OPL, Kaizen, Suggestion etc.,)')
        ]
    },
    'tire_curing': {
        'formatNo': 'Format No: ATC/T/FOR/HR/88D',
        'title': 'ON THE JOB TRAINING EVALUATION - TIRE CURING QA',
        'checkpoints': [
            (1, 'Safety Awareness (BBs, PPE, Tire handling, Electrical safety & VCL )'),
            (2, 'Verification of Press parameters dome temperature, Pressure gauges, Tower lamp working condition , Bladder and sleeve height measurement (Equipment calibrations)'),
            (3, 'Verification on – NSNL /MES/Scanning and recipe'),
            (4, 'Verification of tire engraving covered as per route card'),
            (5, 'Cure cycle time verification as per specification'),
            (6, 'Verification of GT condition (Paint aging and application (Inner & outer)/ GT storage)'),
            (7, 'GT loading direction against route card'),
            (8, 'PCI machine pressure & flange width measurement (OD setting if applicable)'),
            (9, 'NC material handling'),
            (10, 'Knowledge of Measuring equipments Vernier caliper, Measuring tape, Steel rule, dial gauge and Lux meter.')
        ]
    },
    'warehouse': {
        'formatNo': 'Format No: ATC/T/FOR/HR/89D',
        'title': 'ON THE JOB TRAINING EVALUATION - WAREHOUSE QA',
        'checkpoints': [
            (1, 'Safety Awareness (BBs, PPE, Tire handling, Electrical safety,)'),
            (2, 'Basic 5S on shop floor.'),
            (3, 'Precautions for forklift usage during tire loading inside the container.'),
            (4, 'OK tires and not ok tire identification and disposal.'),
            (5, 'Confirmation of PDI cleared tires.'),
            (6, 'Containment action and corrective action for customer complaints/feedbacks.'),
            (7, 'Aging requirements for outgoing product.'),
            (8, 'Poke Yoke in warehouse.'),
            (9, 'Purpose of bead vent/flash trimming on tubeless tires before dispatch.'),
            (10, 'TEI ( Knowledge of QCC, 5S, OPL, Kaizen, Suggestion etc.,)')
        ]
    },
    'solid_tire': {
        'formatNo': 'Format No: ATC/T/FOR/HR/86D',
        'title': 'ON THE JOB TRAINING EVALUATION - SOLID TIRE QA',
        'checkpoints': [
            (1, 'Safety Awareness (BBs, PPE, Tire handling, Electrical safety & VCL)'),
            (2, 'Machine cleanliness and material handling'),
            (3, 'Band building parameter verification (Width, gauge, length & diameter)'),
            (4, 'Base & Tread extrusion profile and temperature verification'),
            (5, 'Press temperature and cure cycle verification'),
            (6, 'Visual inspection of cured solid tires & defect classification'),
            (7, 'Trimming and buffing quality check'),
            (8, 'Pokayoke verification in solid tire press & assembly'),
            (9, 'NC material handling & scrap segregation'),
            (10, 'TEI (Knowledge of QCC, 5S, OPL, Kaizen, Suggestion etc.)')
        ]
    },
    'preparatory': {
        'formatNo': 'Format No: ATC/T/FOR/HR/85D',
        'title': 'ON THE JOB TRAINING EVALUATION - PREPARATORY QA',
        'checkpoints': [
            (1, 'Safety Awareness (BBs, PPE, Material handling, Emergency stops)'),
            (2, 'Raw material and compound receipt verification & FIFO maintenance'),
            (3, 'Calender line parameter check (Fabric tension, cord count, gum gauge)'),
            (4, 'Extruder temperature, screw speed & profile dimension control'),
            (5, 'Bead winding and apexing inspection (Bead diameter, apex height)'),
            (6, 'Slitting and cutting angle & width specification check'),
            (7, 'Storage condition & book aging tracking for rubber components'),
            (8, 'Handling and calibration of measuring instruments (Thickness gauge, vernier)'),
            (9, 'Non-conforming (NC) material identification & tagging'),
            (10, 'TEI (Knowledge of QCC, 5S, OPL, Kaizen, Suggestion etc.)')
        ]
    },
    'final_finish': {
        'formatNo': 'Format No: ATC/T/FOR/HR/84D',
        'title': 'ON THE JOB TRAINING EVALUATION - FINAL FINISH REPAIR ASSOCIATE',
        'checkpoints': [
            (1, 'Safety Awareness (BBs, PPE, Tire handling, Air tool safety)'),
            (2, 'Visual inspection techniques for surface defects (Bead, sidewall, tread)'),
            (3, 'Tire repair criteria assessment (Reparable vs Scrap criteria)'),
            (4, 'Buffing technique, tool selection and depth control'),
            (5, 'Chemical cement application and drying time control'),
            (6, 'Patch application, stitching, and curing parameter check'),
            (7, 'Post-repair inspection & quality sign-off')
        ]
    },
    'rro_alt': {
        'formatNo': 'Format No: ATC/T/FOR/HR/83D',
        'title': 'ON THE JOB TRAINING EVALUATION - RRO & ALT OPERATOR',
        'checkpoints': [
            (1, 'Safety Awareness (BBs, PPE, Rotating machinery safety, Ergonomics)'),
            (2, 'Radial Runout (RRO) & Lateral Runout (LRO) measurement principles'),
            (3, 'Uniformity testing machine calibration & master tire verification'),
            (4, 'Inflation pressure settings and bead seating verification'),
            (5, 'Dynamic balancing procedure and weight placement accuracy'),
            (6, 'High-speed anomaly and force variation detection'),
            (7, 'Classification of tires (Grade A, Grade B, Re-check, Scrap)'),
            (8, 'Marking and data logging in MES/Quality portal')
        ]
    }
}

def normalize_section(sec):
    s = (sec or '').lower().strip()
    s = re.sub(r'\s+', ' ', s)
    s = s.replace('ware house', 'warehouse')
    if 'rro' in s or 'alt' in s: return 'final finish rro & alt qa'
    if 'building' in s or 'tbm' in s: return 'tire building qa'
    if 'curing' in s: return 'tire curing qa'
    if 'solid' in s: return 'solid tire qa'
    if 'preparatory' in s: return 'preparatory qa'
    if 'fid' in s: return 'fid inspector qa'
    if 'warehouse' in s or 'data entry' in s: return 'warehouse qa'
    if 'finish' in s: return 'final finish qa'
    return s

def get_ojt_key(sec_norm):
    if 'rro' in sec_norm or 'alt' in sec_norm: return 'rro_alt'
    if 'preparatory' in sec_norm: return 'preparatory'
    if 'solid' in sec_norm: return 'solid_tire'
    if 'building' in sec_norm: return 'tire_building'
    if 'curing' in sec_norm: return 'tire_curing'
    if 'warehouse' in sec_norm or 'fid' in sec_norm: return 'warehouse'
    return 'final_finish'

def load_data():
    employees = []
    question_bank = {'L': [], 'U': [], 'O': []}
    
    if os.path.exists(DATA_JS):
        with open(DATA_JS, 'r', encoding='utf-8') as f:
            content = f.read()
        emp_m = re.search(r"const EMPLOYEES\s*=\s*(\[.*?\]);", content, re.DOTALL)
        if emp_m:
            try: employees = json.loads(emp_m.group(1))
            except: pass
        qb_m = re.search(r"const QUESTION_BANK\s*=\s*(\{.*?\});\s*(?:const|let|var|function)", content, re.DOTALL)
        if qb_m:
            try: question_bank = json.loads(qb_m.group(1))
            except: pass
            
    ojt_records = {}
    if os.path.exists(OJT_JSON):
        try:
            with open(OJT_JSON, 'r', encoding='utf-8') as f:
                ojt_records = json.load(f)
        except: pass

    assessment_records = {}
    if os.path.exists(RECORDS_JSON):
        try:
            with open(RECORDS_JSON, 'r', encoding='utf-8') as f:
                assessment_records = json.load(f)
        except: pass
            
    return employees, question_bank, ojt_records, assessment_records

def set_cell_value(cell, text, bold=True, color=None, size_pt=9, bg_color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text))
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    if bg_color:
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

def is_option_paragraph(txt):
    t_low = txt.lower().strip()
    if re.match(r'^[a-dA-D][\.\)]', txt): return True
    if t_low.startswith("none of") or t_low.startswith("all the above") or t_low.startswith("all of the above"): return True
    if "மேலே சொன்னது எதுவுமே இல்லை" in txt or "மேலே சொன்னது அனைத்தும்" in txt or "மேலே கூறப்பட்ட அனைத்தும்" in txt: return True
    return False

def is_question_paragraph(txt):
    if is_option_paragraph(txt): return False
    if "Assessment Questionnaire" in txt or "YOKOHAMA OFF-HIGHWAY" in txt: return False
    if txt in ["Safety", "CI & TPM", "QA & Process"]: return False
    if "?" in txt: return True
    if "என்றால் என்ன" in txt or "கண்டறியவும்" in txt or "என்ன?" in txt: return True
    return False

def clean_text_for_match(t):
    return re.sub(r'[\s\.\?\/,:;\(\)]+', '', (t or '').lower())

def generate_employee_docx(emp_no, records_map=None, output_path=None):
    employees, question_bank, ojt_records, file_records = load_data()
    emp = next((e for e in employees if str(e.get('empNo')) == str(emp_no)), None)
    
    if not emp:
        emp = {
            'empNo': emp_no,
            'name': f'Employee {emp_no}',
            'dept': 'QUALITY CONTROL',
            'section': 'Tire building QA',
            'currentLevel': 'O',
            'doj': '-'
        }

    rec = {}
    if records_map and emp_no in records_map:
        rec = records_map[emp_no]
    elif file_records and emp_no in file_records:
        rec = file_records[emp_no]

    ojt = ojt_records.get(emp_no, {})

    curr_lvl = emp.get('currentLevel') or 'I'
    target_map = {'I': 'L', 'L': 'U', 'U': 'O', 'O': 'O'}
    target_lvl = rec.get('targetLevel') or target_map.get(curr_lvl, 'L')
    
    norm_sec = normalize_section(emp.get('section'))
    template_filename = TEMPLATE_MAP.get((target_lvl, norm_sec))
    
    if not template_filename:
        template_filename = f"{target_lvl} Level - Tire Building QA.docx"
        
    template_path = os.path.join(QC_DIR, template_filename)
    if not os.path.exists(template_path):
        cand = [f for f in os.listdir(QC_DIR) if f.startswith(f"{target_lvl} Level") and f.endswith('.docx')]
        if cand: template_path = os.path.join(QC_DIR, cand[0])
        else: raise FileNotFoundError(f"Template not found for level {target_lvl} and section {norm_sec}")

    print(f"Loading template: {os.path.basename(template_path)}")
    doc = docx.Document(template_path)

    # 1. Insert Yokohama Logo & Official Banner at the top
    p_first = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    
    p_logo = p_first.insert_paragraph_before()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(4)
    if os.path.exists(LOGO_PATH):
        r_logo = p_logo.add_run()
        r_logo.add_picture(LOGO_PATH, width=Inches(3.2))

    p_sub = p_first.insert_paragraph_before()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(8)
    r_sub1 = p_sub.add_run("YOKOHAMA OFF-HIGHWAY TIRES (ATC TIRES PVT. LTD.)\n")
    r_sub1.font.name = "Arial"
    r_sub1.font.size = Pt(10)
    r_sub1.font.bold = True
    r_sub1.font.color.rgb = COLOR_NAVY
    
    r_sub2 = p_sub.add_run("QUALITY ASSURANCE DIVISION — ILUO SKILL LEVEL QUALIFICATION RECORD")
    r_sub2.font.name = "Arial"
    r_sub2.font.size = Pt(8.5)
    r_sub2.font.bold = True
    r_sub2.font.color.rgb = COLOR_MUTED

    p_first.insert_paragraph_before()

    # 2. Extract Questions & Marks
    submitted_qs = rec.get('submittedQuestions', [])
    is_attempted = bool(rec.get('isCompleted') or rec.get('inProgress') or submitted_qs)
    
    total_mark = rec.get('totalMark') if rec.get('totalMark') is not None else (len([q for q in submitted_qs if q.get('isCorrect')]) if submitted_qs else 0)
    total_possible = len(submitted_qs) if submitted_qs else 30
    min_pass = int(total_possible * 0.7 + 0.99)
    is_passed = (rec.get('status') == 'Passed') or (total_mark >= min_pass and is_attempted)
    
    has_ojt = bool(ojt and (ojt.get('totalScore') is not None))
    ojt_score = ojt.get('totalScore', 0) if has_ojt else 0
    ojt_max = ojt.get('maxScore', 50) if has_ojt else 50
    ojt_qualified = bool(has_ojt and (ojt.get('qualificationStatus') == 'Qualified' or ojt_score >= (ojt_max * 0.7)))

    if is_passed and ojt_qualified:
        status_str = f"QUALIFIED — LEVEL UP TO {target_lvl}"
        status_color = COLOR_GREEN
    elif is_passed and not has_ojt:
        status_str = "THEORY PASSED (OJT PENDING)"
        status_color = COLOR_YOKO_BLUE
    elif not is_passed and is_attempted and has_ojt and ojt_qualified:
        status_str = "OJT QUALIFIED (THEORY RETEST)"
        status_color = RGBColor(180, 83, 9)
    elif is_attempted and not is_passed:
        status_str = "RETEST REQUIRED (MIN. MARKS NOT MET)"
        status_color = COLOR_RED
    else:
        status_str = "ASSESSMENT PENDING"
        status_color = COLOR_MUTED

    # 3. Fill Table 0 (Metadata Table)
    if doc.tables:
        t = doc.tables[0]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        emp_name = emp.get('name') or rec.get('name') or emp_no
        emp_dept = emp.get('dept') or rec.get('dept') or 'QUALITY CONTROL'
        emp_sec = emp.get('section') or rec.get('section') or '-'
        emp_doj = emp.get('doj') or rec.get('doj') or '-'
        exam_date = rec.get('attemptDate') or datetime.now().strftime('%d/%m/%Y')
        
        # Row 0: NAME | EMPLOYEE NO
        set_cell_value(t.cell(0, 1), emp_name, bold=True, color=COLOR_DARK)
        set_cell_value(t.cell(0, 3), emp_no, bold=True, color=COLOR_YOKO_BLUE)
        
        # Row 1: DEPARTMENT | DOJ
        set_cell_value(t.cell(1, 1), emp_dept, bold=False, color=COLOR_DARK)
        set_cell_value(t.cell(1, 3), emp_doj, bold=False, color=COLOR_DARK)
        
        # Row 2: SECTION | DATE
        set_cell_value(t.cell(2, 1), emp_sec, bold=True, color=COLOR_YOKO_BLUE)
        set_cell_value(t.cell(2, 3), exam_date, bold=False, color=COLOR_DARK)
        
        # Row 3: TOTAL MARKS | RESULT (Strictly Marks Only, Zero %)
        marks_display = f"{total_mark} / {total_possible} Marks" if is_attempted else "Pending Exam"
        set_cell_value(t.cell(3, 1), marks_display, bold=True, color=COLOR_GREEN if is_passed else (COLOR_RED if is_attempted else COLOR_MUTED))
        
        t.cell(3, 2).text = "STATUS / RESULT"
        t.cell(3, 2).paragraphs[0].runs[0].font.bold = True
        t.cell(3, 2).paragraphs[0].runs[0].font.size = Pt(9)
        set_cell_value(t.cell(3, 3), status_str, bold=True, color=status_color)

    # 4. Map and Mark Questions in the Document
    q_by_idx = {q.get('index', i+1): q for i, q in enumerate(submitted_qs)}
    q_by_text = {clean_text_for_match(q.get('question', ''))[:30]: q for q in submitted_qs if q.get('question')}

    # Pre-parse questions in document to assign sequential question indices
    doc_questions = []
    current_q_obj = None

    for idx, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if not txt: continue

        if is_question_paragraph(txt):
            current_q_obj = {
                'p_idx': idx,
                'q_p': p,
                'q_text': txt,
                'options': []
            }
            doc_questions.append(current_q_obj)
        elif current_q_obj is not None:
            if not ("Assessment Questionnaire" in txt or txt in ["Safety", "CI & TPM", "QA & Process"]):
                current_q_obj['options'].append((idx, p, txt))

    # Annotate parsed questions
    for q_idx_zero, q_entry in enumerate(doc_questions):
        q_num = q_idx_zero + 1
        matched_q = q_by_idx.get(q_num)
        if not matched_q:
            q_clean = clean_text_for_match(q_entry['q_text'])[:30]
            matched_q = q_by_text.get(q_clean)

        if not matched_q:
            continue

        is_corr = matched_q.get('isCorrect', False)
        sel_key = (matched_q.get('selectedKey') or '').strip().upper()
        corr_key = (matched_q.get('correctKey') or '').strip().upper()

        # Add mark badge to question paragraph
        q_p = q_entry['q_p']
        r_mark = q_p.add_run(f"   [Mark: 1/1 — Correct]" if is_corr else f"   [Mark: 0/1 — Incorrect]")
        r_mark.font.bold = True
        r_mark.font.size = Pt(8.5)
        r_mark.font.color.rgb = COLOR_GREEN if is_corr else COLOR_RED

        # Case A: If question had all options in a single paragraph with \n
        if "\n" in q_entry['q_text'] and ("a." in q_entry['q_text'] or "a )" in q_entry['q_text'] or "a)" in q_entry['q_text']):
            lines = q_entry['q_text'].split("\n")
            q_p.text = ""
            r_q0 = q_p.add_run(lines[0] + "   ")
            r_q0.font.bold = True
            r_m = q_p.add_run(f"[Mark: 1/1 — Correct]\n" if is_corr else f"[Mark: 0/1 — Incorrect]\n")
            r_m.font.bold = True
            r_m.font.size = Pt(8.5)
            r_m.font.color.rgb = COLOR_GREEN if is_corr else COLOR_RED

            for line in lines[1:]:
                line_str = line.strip()
                km = re.match(r'^([a-dA-D])[\.\)]', line_str)
                k = km.group(1).upper() if km else ""
                
                if k and k == sel_key:
                    if is_corr:
                        r_opt = q_p.add_run(f"  [✔ CANDIDATE SELECTED - CORRECT (+1 Mark)]  {line_str}\n")
                        r_opt.font.bold = True
                        r_opt.font.color.rgb = COLOR_GREEN
                    else:
                        r_opt = q_p.add_run(f"  [✘ CANDIDATE SELECTED - INCORRECT (0 Marks)]  {line_str}\n")
                        r_opt.font.bold = True
                        r_opt.font.color.rgb = COLOR_RED
                elif k and k == corr_key and not is_corr:
                    r_opt = q_p.add_run(f"  [✔ OFFICIAL ANSWER KEY]  {line_str}\n")
                    r_opt.font.bold = True
                    r_opt.font.color.rgb = COLOR_GREEN
                else:
                    r_opt = q_p.add_run(f"  {line_str}\n")
                    r_opt.font.color.rgb = COLOR_DARK
            continue

        # Case B: Multi-paragraph options
        key_chars = ['A', 'B', 'C', 'D']
        for opt_idx, (_, opt_p, opt_txt) in enumerate(q_entry['options'][:4]):
            opt_key_m = re.match(r'^([a-dA-D])[\.\)]', opt_txt)
            inferred_key = opt_key_m.group(1).upper() if opt_key_m else key_chars[opt_idx]

            is_sel = (inferred_key == sel_key)
            is_key = (inferred_key == corr_key)

            if is_sel:
                opt_p.text = ""
                if is_corr:
                    r_o = opt_p.add_run(f"[✔ CANDIDATE SELECTED - CORRECT (+1 Mark)]  {opt_txt}")
                    r_o.font.bold = True
                    r_o.font.color.rgb = COLOR_GREEN
                else:
                    r_o = opt_p.add_run(f"[✘ CANDIDATE SELECTED - INCORRECT (0 Marks)]  {opt_txt}")
                    r_o.font.bold = True
                    r_o.font.color.rgb = COLOR_RED
            elif is_key and not is_corr:
                opt_p.text = ""
                r_o = opt_p.add_run(f"[✔ OFFICIAL ANSWER KEY]  {opt_txt}")
                r_o.font.bold = True
                r_o.font.color.rgb = COLOR_GREEN

    # 5. Append Practical OJT Evaluation Section
    p_ojt_head = doc.add_paragraph()
    p_ojt_head.paragraph_format.space_before = Pt(14)
    p_ojt_head.paragraph_format.space_after = Pt(4)
    r_oh = p_ojt_head.add_run("PRACTICAL ON-THE-JOB TRAINING (OJT) EVALUATION REPORT")
    r_oh.font.name = "Arial"
    r_oh.font.size = Pt(11)
    r_oh.font.bold = True
    r_oh.font.color.rgb = COLOR_NAVY

    ojt_info = OJT_CHECKPOINTS.get(get_ojt_key(norm_sec), OJT_CHECKPOINTS['tire_building'])

    if has_ojt:
        ojt_scores = ojt.get('scores', {})
        t_ojt = doc.add_table(rows=1, cols=3)
        t_ojt.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        hdr_cells = t_ojt.rows[0].cells
        hdr_cells[0].text = "S.No"
        hdr_cells[1].text = "Practical Evaluation Checkpoint"
        hdr_cells[2].text = "Marks Awarded"
        
        for c in hdr_cells:
            c.paragraphs[0].runs[0].font.bold = True
            c.paragraphs[0].runs[0].font.size = Pt(9)
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
            c._tc.get_or_add_tcPr().append(shading)
            
        hdr_cells[0].width = Inches(0.6)
        hdr_cells[1].width = Inches(4.5)
        hdr_cells[2].width = Inches(1.5)

        for s_no, cp_desc in ojt_info['checkpoints']:
            sc = ojt_scores.get(str(s_no), ojt_scores.get(s_no, '-'))
            sc_str = f"{sc} / 5 Marks" if sc != '-' else "- / 5 Marks"
            row_cells = t_ojt.add_row().cells
            set_cell_value(row_cells[0], str(s_no), bold=True, color=COLOR_MUTED)
            set_cell_value(row_cells[1], cp_desc, bold=False, color=COLOR_DARK)
            set_cell_value(row_cells[2], sc_str, bold=True, color=COLOR_YOKO_BLUE)

        p_comm = doc.add_paragraph()
        p_comm.paragraph_format.space_before = Pt(6)
        r_c1 = p_comm.add_run(f"Practical Score: {ojt_score} / {ojt_max} Marks  •  Status: {ojt.get('qualificationStatus', 'Qualified')}\n")
        r_c1.font.bold = True
        r_c1.font.size = Pt(9.5)
        r_c1.font.color.rgb = COLOR_GREEN if ojt_qualified else COLOR_RED
        
        r_c2 = p_comm.add_run(f"Improvement Comments: {ojt.get('comments') or 'Practical evaluation satisfactory.'}\n")
        r_c2.font.size = Pt(9)
        r_c3 = p_comm.add_run(f"Evaluation Committee: Safety: {ojt.get('safetyRep', '-')} | Quality: {ojt.get('qualityRep', '-')} | CI: {ojt.get('ciRep', '-')}")
        r_c3.font.size = Pt(8.5)
        r_c3.font.color.rgb = COLOR_MUTED
    else:
        p_pend = doc.add_paragraph()
        p_pend.paragraph_format.space_before = Pt(4)
        r_p = p_pend.add_run("Practical OJT Status: In-section practical evaluation is pending for this employee. Practical marks will append upon supervisor submission.")
        r_p.font.size = Pt(9)
        r_p.font.italic = True
        r_p.font.color.rgb = COLOR_MUTED

    # 6. Official Sign-Off Certification Block
    p_sig_head = doc.add_paragraph()
    p_sig_head.paragraph_format.space_before = Pt(16)
    
    t_sig = doc.add_table(rows=2, cols=3)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_widths = [Inches(2.2), Inches(2.2), Inches(2.2)]
    
    for c_idx, cell in enumerate(t_sig.rows[0].cells):
        cell.width = sig_widths[c_idx]
        set_cell_value(cell, "\n\n___________________________", bold=True, color=COLOR_MUTED)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    labels = [
        f"Candidate Signature\n{emp_name} ({emp_no})",
        "QA Section Evaluator\nTechnical Incharge / Supervisor",
        "Quality Assurance Manager\nPlant Quality Head Approval & Seal"
    ]
    for c_idx, cell in enumerate(t_sig.rows[1].cells):
        cell.width = sig_widths[c_idx]
        set_cell_value(cell, labels[c_idx], bold=True, color=COLOR_DARK, size_pt=8.5)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(10)
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_f = p_foot.add_run(f"YOKOHAMA OFF-HIGHWAY TIRES • OFFICIAL QA ILUO AUDIT RECORD • STRICTLY CONFIDENTIAL\nGenerated: {datetime.now().strftime('%d/%m/%Y, %H:%M:%S')}")
    r_f.font.size = Pt(7.5)
    r_f.font.color.rgb = COLOR_MUTED

    safe_name = re.sub(r'[\s\\/]+', '_', emp_name)
    if not output_path:
        output_path = os.path.join(OUTPUT_DIR, f"Yokohama_ILUO_Report_{emp_no}_{safe_name}.docx")
        
    doc.save(output_path)
    print(f"SUCCESS: Generated DOCX for Employee {emp_no} -> {output_path}")
    return output_path

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate Yokohama Employee ILUO Evaluation DOCX Report")
    parser.add_argument("--emp", type=str, help="Employee ID (e.g. 900262)")
    parser.add_argument("--output", type=str, help="Output DOCX file path")
    parser.add_argument("--record-file", type=str, help="Path to JSON file containing employee assessment record")
    parser.add_argument("--all", action="store_true", help="Generate DOCX for all employees")
    args = parser.parse_args()

    records_map = None
    if args.record_file and os.path.exists(args.record_file):
        try:
            with open(args.record_file, 'r', encoding='utf-8') as f:
                rec_data = json.load(f)
                if args.emp and args.emp in rec_data:
                    records_map = {args.emp: rec_data[args.emp]}
                elif args.emp:
                    records_map = {args.emp: rec_data}
                else:
                    records_map = rec_data
        except Exception as e:
            print(f"Warning: Could not read record file: {e}")

    if args.emp:
        generate_employee_docx(args.emp, records_map=records_map, output_path=args.output)
    elif args.all:
        employees, _, _, _ = load_data()
        print(f"Generating DOCX for all {len(employees)} employees...")
        count = 0
        for emp in employees:
            try:
                generate_employee_docx(emp['empNo'], records_map=records_map)
                count += 1
            except Exception as e:
                print(f"Error generating for {emp.get('empNo')}: {e}")
        print(f"Finished generating {count} DOCX files in {OUTPUT_DIR}")
    else:
        print("Usage: python generate_employee_docx.py --emp <empNo> [--output <path>] [--record-file <file>] or --all")

