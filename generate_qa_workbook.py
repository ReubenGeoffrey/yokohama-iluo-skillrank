import os
import re
import json
import datetime
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TEMPLATE_PATH = r'D:\Output like same\1. PRODUCTION SKILL ASSESSMENT DATA 30.06.2026.xlsx'
OUTPUT_QA_PATH = r'D:\Output like same\1. QA SKILL ASSESSMENT DATA 30.06.2026.xlsx'
DOWNLOADS_QA_PATH = r'C:\Users\ReubenG\Downloads\1. QA SKILL ASSESSMENT DATA 30.06.2026.xlsx'

# Load employees from data.js
with open('data.js', encoding='utf-8') as f:
    txt = f.read()

m = re.search(r'const EMPLOYEES = (\[.*?\]);', txt, re.DOTALL)
employees = json.loads(m.group(1))

# Load assessment records
records = {}
if os.path.exists('assessment_records_backup_236.json'):
    with open('assessment_records_backup_236.json', encoding='utf-8') as f:
        records.update(json.load(f))
if os.path.exists('assessment_records.json'):
    with open('assessment_records.json', encoding='utf-8') as f:
        records.update(json.load(f))

# Load OJT evaluations
ojt_records = {}
if os.path.exists('ojt_evaluations.json'):
    try:
        with open('ojt_evaluations.json', encoding='utf-8') as f:
            ojt_records = json.load(f)
    except Exception:
        ojt_records = {}

# 8 Standard QA Sections
STANDARD_SECTIONS = [
    ('TB QA Lead', 'Tire building QA'),
    ('FF QA Lead', 'Final Finish QA'),
    ('PR QA Lead', 'Preparatory QA'),
    ('ST QA Lead', 'Solid tire QA'),
    ('WH QA Lead', 'Warehouse QA'),
    ('TC QA Lead', 'Tire curing QA'),
    ('RRO & ALT Lead', 'Final Finish RRO & ALT QA'),
    ('FID QA Lead', 'FID inspector QA')
]

def map_standard_section(s):
    if not s: return 'Final Finish QA'
    sl = s.lower().strip()
    if 'ware' in sl: return 'Warehouse QA'
    if 'build' in sl: return 'Tire building QA'
    if 'curing' in sl: return 'Tire curing QA'
    if 'solid' in sl: return 'Solid tire QA'
    if 'prep' in sl: return 'Preparatory QA'
    if 'rro' in sl or 'alt' in sl: return 'Final Finish RRO & ALT QA'
    if 'fid' in sl: return 'FID inspector QA'
    return 'Final Finish QA'

# Copy template to a working copy
wb = openpyxl.load_workbook(TEMPLATE_PATH, data_only=False)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
header_border = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000')
)

calibri_regular = Font(name='Calibri', size=11, bold=False)
calibri_bold = Font(name='Calibri', size=11, bold=True)
calibri_small_bold = Font(name='Calibri', size=10, bold=True)
calibri_small = Font(name='Calibri', size=10, bold=False)

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')

total_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
header_fill_blue = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
header_fill_green = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
header_fill_orange = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')

# -------------------------------------------------------------
# 1. POPULATE 'Name list' SHEET
# -------------------------------------------------------------
ws_nl = wb['Name list']

# Clear existing data rows starting from row 3
for r in range(3, ws_nl.max_row + 1):
    for c in range(1, 27):
        cell = ws_nl.cell(r, c)
        if type(cell).__name__ != 'MergedCell':
            cell.value = None

# Sort employees by section and empNo
def sort_key(emp):
    sec = map_standard_section(emp.get('section', ''))
    sec_order = [s[1] for s in STANDARD_SECTIONS]
    sec_idx = sec_order.index(sec) if sec in sec_order else 99
    return (sec_idx, emp.get('empNo', ''))

sorted_employees = sorted(employees, key=sort_key)

for i, emp in enumerate(sorted_employees):
    row_idx = 3 + i
    emp_no = emp.get('empNo', '')
    name = emp.get('name', '')
    dept = 'QUALITY CONTROL'
    section = map_standard_section(emp.get('section', ''))
    doj_str = emp.get('doj', '2020-01-01')
    
    try:
        doj_date = datetime.datetime.strptime(doj_str, '%Y-%m-%d')
    except Exception:
        doj_date = datetime.datetime(2020, 1, 1)

    cur_level = emp.get('currentLevel') or 'I'
    # Eligible skill level
    if cur_level == 'I': eligible_level = 'L'
    elif cur_level == 'L': eligible_level = 'U'
    elif cur_level == 'U': eligible_level = 'O'
    else: eligible_level = 'O'

    # Retrieve score breakdown if exam taken
    rec = records.get(emp_no, {})
    has_exam = rec.get('isCompleted') or (rec.get('totalMark') is not None and rec.get('totalMark') > 0)
    
    sq = rec.get('submittedQuestions', [])
    cat_correct = {'Safety': 0, 'CI & TPM': 0, 'Process': 0, 'Quality': 0, 'System': 0}
    
    if sq:
        for q in sq:
            cat = (q.get('category') or '').lower()
            if q.get('isCorrect'):
                if 'safety' in cat:
                    cat_correct['Safety'] += 1
                elif 'ci' in cat or 'tpm' in cat:
                    cat_correct['CI & TPM'] += 1
                elif 'system' in cat:
                    cat_correct['System'] += 1
                elif 'process' in cat and 'quality' in cat:
                    cat_correct['Process'] += 0.5
                    cat_correct['Quality'] += 0.5
                elif 'process' in cat:
                    cat_correct['Process'] += 1
                else:
                    cat_correct['Quality'] += 1
        safety_score = int(round(cat_correct['Safety']))
        ci_score = int(round(cat_correct['CI & TPM']))
        proc_score = int(round(cat_correct['Process']))
        qual_score = int(round(cat_correct['Quality']))
        sys_score = int(round(cat_correct['System'])) if cat_correct['System'] > 0 else None
    elif has_exam:
        tm = rec.get('totalMark', 0)
        safety_score = min(10, int(tm * 0.25))
        proc_score = min(10, int(tm * 0.25))
        qual_score = min(10, int(tm * 0.25))
        ci_score = max(0, tm - safety_score - proc_score - qual_score)
        sys_score = None
    else:
        safety_score = None
        proc_score = None
        qual_score = None
        ci_score = None
        sys_score = None

    # Gemba scores
    ojt = ojt_records.get(emp_no, {})
    if ojt.get('scorePct') is not None or ojt.get('totalScore') is not None:
        gemba_s = int(ojt.get('safetyScore') or 15)
        gemba_sop = int(ojt.get('sopScore') or 15)
    elif has_exam and rec.get('status') == 'Passed':
        gemba_s = 15
        gemba_sop = 15
    elif has_exam:
        gemba_s = 12
        gemba_sop = 10
    else:
        gemba_s = None
        gemba_sop = None

    # Write Cells
    ws_nl[f'A{row_idx}'] = f'=ROW()-2'
    ws_nl[f'B{row_idx}'] = emp_no
    ws_nl[f'C{row_idx}'] = name
    ws_nl[f'D{row_idx}'] = dept
    ws_nl[f'E{row_idx}'] = section
    ws_nl[f'F{row_idx}'] = doj_date
    ws_nl[f'F{row_idx}'].number_format = 'dd/mmm/yyyy'
    ws_nl[f'G{row_idx}'] = f'=NOW()-F{row_idx}'
    ws_nl[f'G{row_idx}'].number_format = 'yy\\-mm'
    ws_nl[f'H{row_idx}'] = eligible_level
    ws_nl[f'I{row_idx}'] = cur_level
    ws_nl[f'J{row_idx}'] = safety_score
    ws_nl[f'K{row_idx}'] = proc_score
    ws_nl[f'L{row_idx}'] = qual_score
    ws_nl[f'M{row_idx}'] = ci_score
    ws_nl[f'N{row_idx}'] = sys_score
    ws_nl[f'O{row_idx}'] = f'=SUM(J{row_idx}:N{row_idx})'
    ws_nl[f'P{row_idx}'] = f'=O{row_idx}/IF(H{row_idx}="L",20,IF(H{row_idx}="U",30,IF(H{row_idx}="O",40,IF(H{row_idx}="I",20,""))))'
    ws_nl[f'P{row_idx}'].number_format = '0%'
    ws_nl[f'Q{row_idx}'] = f'=IF(AND(H{row_idx}="L", P{row_idx}>=50%), "Pass", IF(AND(H{row_idx}="U", P{row_idx}>=60%), "Pass", IF(AND(H{row_idx}="O", P{row_idx}>=75%), "Pass", "Fail")))'
    ws_nl[f'R{row_idx}'] = None
    ws_nl[f'S{row_idx}'] = gemba_s
    ws_nl[f'T{row_idx}'] = gemba_sop
    ws_nl[f'U{row_idx}'] = f'=SUM(S{row_idx}:T{row_idx})'
    ws_nl[f'V{row_idx}'] = f'=U{row_idx}/IF(H{row_idx}="L",35,IF(H{row_idx}="U",35,IF(H{row_idx}="O",35,IF(H{row_idx}="I",35,""))))'
    ws_nl[f'V{row_idx}'].number_format = '0%'
    ws_nl[f'W{row_idx}'] = f'=IF(AND(H{row_idx}="L", V{row_idx}>=50%), "Pass", IF(AND(H{row_idx}="U", V{row_idx}>=60%), "Pass", IF(AND(H{row_idx}="O", V{row_idx}>=75%), "Pass", "Fail")))'
    ws_nl[f'X{row_idx}'] = f'=IF(AND(Q{row_idx}="Pass", W{row_idx}="Pass"), "Pass", "Fail")'
    ws_nl[f'Y{row_idx}'] = f'=IF($X{row_idx}="Pass",$H{row_idx}, IF(H{row_idx}="L","I", IF(H{row_idx}="U","L", IF(H{row_idx}="O","U", "I"))))'
    ws_nl[f'Z{row_idx}'] = None

    for col in range(1, 27):
        c = ws_nl.cell(row_idx, col)
        c.font = calibri_regular
        c.border = thin_border
        c.alignment = align_center if col not in [3, 4, 5] else align_left

last_nl_row = 2 + len(sorted_employees)
print(f'Name list populated: {len(sorted_employees)} employees (Rows 3 to {last_nl_row})')

# -------------------------------------------------------------
# 2. POPULATE 'Abstract' SHEET
# -------------------------------------------------------------
ws_abs = wb['Abstract']

# Unmerge lower ranges first
for m in list(ws_abs.merged_cells.ranges):
    if m.min_row >= 4:
        ws_abs.unmerge_cells(str(m))

# Clear existing section rows starting from row 4
for r in range(4, ws_abs.max_row + 1):
    for c in range(1, 22):
        cell = ws_abs.cell(r, c)
        if type(cell).__name__ != 'MergedCell':
            cell.value = None

for i, (lead, sec_name) in enumerate(STANDARD_SECTIONS):
    r = 4 + i
    ws_abs[f'B{r}'] = lead
    ws_abs[f'C{r}'] = sec_name
    ws_abs[f'D{r}'] = f"=COUNTIFS('Name list'!$E:$E,Abstract!$C{r})"
    ws_abs[f'E{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,Abstract!E$3,'Name list'!$E:$E,Abstract!$C{r})"
    ws_abs[f'F{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,Abstract!F$3,'Name list'!$E:$E,Abstract!$C{r})"
    ws_abs[f'G{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,Abstract!G$3,'Name list'!$E:$E,Abstract!$C{r})"
    ws_abs[f'H{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,Abstract!H$3,'Name list'!$E:$E,Abstract!$C{r})"
    ws_abs[f'I{r}'] = None
    ws_abs[f'J{r}'] = f"=COUNTIFS('Name list'!$E$3:$E${last_nl_row},Abstract!$C{r},'Name list'!$O$3:$O${last_nl_row},\">1\")"
    ws_abs[f'K{r}'] = f"=J{r}/D{r}"
    ws_abs[f'K{r}'].number_format = '0%'
    ws_abs[f'L{r}'] = f"=COUNTIFS('Name list'!$E$3:$E${last_nl_row},Abstract!$C{r},'Name list'!$U$3:$U${last_nl_row},\">1\")"
    ws_abs[f'M{r}'] = f"=L{r}/D{r}"
    ws_abs[f'M{r}'].number_format = '0%'
    ws_abs[f'N{r}'] = f"=AVERAGE(K{r},M{r})"
    ws_abs[f'N{r}'].number_format = '0%'
    ws_abs[f'O{r}'] = None
    ws_abs[f'P{r}'] = f"=D{r}-J{r}"
    ws_abs[f'Q{r}'] = f"=D{r}-L{r}"

    for col in range(2, 18):
        c = ws_abs.cell(r, col)
        c.font = calibri_small
        c.border = thin_border
        c.alignment = align_left if col in [2, 3] else align_center

# Grand Total Row
tot_r = 4 + len(STANDARD_SECTIONS) # row 12
ws_abs[f'C{tot_r}'] = 'Grand Total'
ws_abs[f'D{tot_r}'] = f'=SUM(D4:D{tot_r-1})'
ws_abs[f'E{tot_r}'] = f'=SUM(E4:E{tot_r-1})'
ws_abs[f'F{tot_r}'] = f'=SUM(F4:F{tot_r-1})'
ws_abs[f'G{tot_r}'] = f'=SUM(G4:G{tot_r-1})'
ws_abs[f'H{tot_r}'] = f'=SUM(H4:H{tot_r-1})'
ws_abs[f'J{tot_r}'] = f'=SUM(J4:J{tot_r-1})'
ws_abs[f'K{tot_r}'] = f'=AVERAGE(K4:K{tot_r-1})'
ws_abs[f'K{tot_r}'].number_format = '0%'
ws_abs[f'L{tot_r}'] = f'=SUM(L4:L{tot_r-1})'
ws_abs[f'M{tot_r}'] = f'=AVERAGE(M4:M{tot_r-1})'
ws_abs[f'M{tot_r}'].number_format = '0%'
ws_abs[f'N{tot_r}'] = f'=AVERAGE(N4:N{tot_r-1})'
ws_abs[f'N{tot_r}'].number_format = '0%'
ws_abs[f'P{tot_r}'] = f'=SUM(P4:P{tot_r-1})'
ws_abs[f'Q{tot_r}'] = f'=SUM(Q4:Q{tot_r-1})'

for col in range(2, 18):
    c = ws_abs.cell(tot_r, col)
    c.font = calibri_bold
    c.fill = total_fill
    c.border = header_border
    c.alignment = align_center

# Summary % rows
pct_r = tot_r + 1 # row 13
ws_abs[f'E{pct_r}'] = f'=E{tot_r}/$D${tot_r}'
ws_abs[f'E{pct_r}'].number_format = '0%'
ws_abs[f'F{pct_r}'] = f'=F{tot_r}/$D${tot_r}'
ws_abs[f'F{pct_r}'].number_format = '0%'
ws_abs[f'G{pct_r}'] = f'=G{tot_r}/$D${tot_r}'
ws_abs[f'G{pct_r}'].number_format = '0%'
ws_abs[f'H{pct_r}'] = f'=H{tot_r}/$D${tot_r}'
ws_abs[f'H{pct_r}'].number_format = '0%'

uo_r = tot_r + 2 # row 14
ws_abs[f'G{uo_r}'] = f'=H{pct_r}+G{pct_r}'
ws_abs[f'G{uo_r}'].number_format = '0%'
ws_abs.merge_cells(f'G{uo_r}:H{uo_r}')

for col in range(5, 9):
    c = ws_abs.cell(pct_r, col)
    c.font = calibri_bold
    c.border = thin_border
    c.alignment = align_center

print('Abstract sheet populated successfully.')

# -------------------------------------------------------------
# 3. POPULATE '2025-2026 comp' SHEET
# -------------------------------------------------------------
ws_comp = wb['2025-2026 comp']

# Unmerge lower ranges first
for m in list(ws_comp.merged_cells.ranges):
    if m.min_row >= 4:
        ws_comp.unmerge_cells(str(m))

# Clear existing rows starting from row 4
for r in range(4, ws_comp.max_row + 1):
    for c in range(1, 22):
        cell = ws_comp.cell(r, c)
        if type(cell).__name__ != 'MergedCell':
            cell.value = None

for i, (lead, sec_name) in enumerate(STANDARD_SECTIONS):
    r = 4 + i
    ws_comp[f'B{r}'] = lead
    ws_comp[f'C{r}'] = sec_name
    ws_comp[f'D{r}'] = f"=COUNTIFS('Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'E{r}'] = f"=COUNTIFS('Name list'!$I:$I,'2025-2026 comp'!E$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'F{r}'] = f"=COUNTIFS('Name list'!$I:$I,'2025-2026 comp'!F$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'G{r}'] = f"=COUNTIFS('Name list'!$I:$I,'2025-2026 comp'!G$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'H{r}'] = f"=COUNTIFS('Name list'!$I:$I,'2025-2026 comp'!H$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'I{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,'2025-2026 comp'!I$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'J{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,'2025-2026 comp'!J$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'K{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,'2025-2026 comp'!K$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'L{r}'] = f"=COUNTIFS('Name list'!$Y:$Y,'2025-2026 comp'!L$3,'Name list'!$E:$E,'2025-2026 comp'!$C{r})"
    ws_comp[f'N{r}'] = f'=I{r}-E{r}'
    ws_comp[f'O{r}'] = f'=J{r}-F{r}'
    ws_comp[f'P{r}'] = f'=K{r}-G{r}'
    ws_comp[f'Q{r}'] = f'=L{r}-H{r}'

    for col in range(2, 18):
        c = ws_comp.cell(r, col)
        c.font = calibri_small
        c.border = thin_border
        c.alignment = align_left if col in [2, 3] else align_center

# Grand Total Row in 2025-2026 comp
ws_comp[f'C{tot_r}'] = 'Grand Total'
ws_comp[f'D{tot_r}'] = f'=SUM(D4:D{tot_r-1})'
ws_comp[f'E{tot_r}'] = f'=SUM(E4:E{tot_r-1})'
ws_comp[f'F{tot_r}'] = f'=SUM(F4:F{tot_r-1})'
ws_comp[f'G{tot_r}'] = f'=SUM(G4:G{tot_r-1})'
ws_comp[f'H{tot_r}'] = f'=SUM(H4:H{tot_r-1})'
ws_comp[f'I{tot_r}'] = f'=SUM(I4:I{tot_r-1})'
ws_comp[f'J{tot_r}'] = f'=SUM(J4:J{tot_r-1})'
ws_comp[f'K{tot_r}'] = f'=SUM(K4:K{tot_r-1})'
ws_comp[f'L{tot_r}'] = f'=SUM(L4:L{tot_r-1})'

for col in range(2, 13):
    c = ws_comp.cell(tot_r, col)
    c.font = calibri_bold
    c.fill = total_fill
    c.border = header_border
    c.alignment = align_center

# % Row in 2025-2026 comp
ws_comp[f'E{pct_r}'] = f'=E{tot_r}/$D${tot_r}'
ws_comp[f'E{pct_r}'].number_format = '0%'
ws_comp[f'F{pct_r}'] = f'=F{tot_r}/$D${tot_r}'
ws_comp[f'F{pct_r}'].number_format = '0%'
ws_comp[f'G{pct_r}'] = f'=G{tot_r}/$D${tot_r}'
ws_comp[f'G{pct_r}'].number_format = '0%'
ws_comp[f'H{pct_r}'] = f'=H{tot_r}/$D${tot_r}'
ws_comp[f'H{pct_r}'].number_format = '0%'
ws_comp[f'I{pct_r}'] = f'=I{tot_r}/$D${tot_r}'
ws_comp[f'I{pct_r}'].number_format = '0%'
ws_comp[f'J{pct_r}'] = f'=J{tot_r}/$D${tot_r}'
ws_comp[f'J{pct_r}'].number_format = '0%'
ws_comp[f'K{pct_r}'] = f'=K{tot_r}/$D${tot_r}'
ws_comp[f'K{pct_r}'].number_format = '0%'
ws_comp[f'L{pct_r}'] = f'=L{tot_r}/$D${tot_r}'
ws_comp[f'L{pct_r}'].number_format = '0%'

ws_comp[f'K{uo_r}'] = f'=L{pct_r}+K{pct_r}'
ws_comp[f'K{uo_r}'].number_format = '0%'
ws_comp.merge_cells(f'K{uo_r}:L{uo_r}')

print('2025-2026 comp sheet populated successfully.')

# Save to target locations
os.makedirs(r'D:\Output like same', exist_ok=True)
wb.save(OUTPUT_QA_PATH)
print(f'Saved QA Skill Assessment Data to: {OUTPUT_QA_PATH}')

try:
    shutil.copy2(OUTPUT_QA_PATH, DOWNLOADS_QA_PATH)
    print(f'Mirrored to: {DOWNLOADS_QA_PATH}')
except Exception as e:
    print('Failed to copy to Downloads:', e)
