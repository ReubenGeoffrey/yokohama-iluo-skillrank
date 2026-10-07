import openpyxl
import datetime
import json
import re
import os

EXCEL_SOURCE = r'D:\Output like same\1. PRODUCTION SKILL ASSESSMENT DATA 30.06.2026.backup.xlsx'

print('Loading production data from:', EXCEL_SOURCE)
wb = openpyxl.load_workbook(EXCEL_SOURCE, data_only=True)
ws = wb['Name list']

# 1. Parse all 1,179 Production employees from rows 3 to 1181
prod_emps = []
prod_records = {}
prod_ojt = {}

for r in range(3, 1182):
    emp_no = str(ws.cell(r, 2).value or '').strip()
    name = str(ws.cell(r, 3).value or '').strip()
    dept = str(ws.cell(r, 4).value or 'PRODUCTION').strip()
    sec = str(ws.cell(r, 5).value or '').strip()
    doj_val = ws.cell(r, 6).value
    if isinstance(doj_val, (datetime.datetime, datetime.date)):
        doj_str = doj_val.strftime('%Y-%m-%d')
    else:
        doj_str = str(doj_val or '2020-01-01')[:10]
    
    elig = str(ws.cell(r, 8).value or 'L').strip()
    cur = str(ws.cell(r, 9).value or 'I').strip()
    
    s_j = ws.cell(r, 10).value
    s_k = ws.cell(r, 11).value
    s_l = ws.cell(r, 12).value
    s_m = ws.cell(r, 13).value
    s_n = ws.cell(r, 14).value
    s_o = ws.cell(r, 15).value
    s_p = ws.cell(r, 16).value
    s_q = ws.cell(r, 17).value
    g_s = ws.cell(r, 19).value
    g_t = ws.cell(r, 20).value
    g_u = ws.cell(r, 21).value
    g_v = ws.cell(r, 22).value
    g_w = ws.cell(r, 23).value
    f_x = ws.cell(r, 24).value
    p_y = ws.cell(r, 25).value

    # Compute experience in months
    try:
        doj_d = datetime.datetime.strptime(doj_str, '%Y-%m-%d')
        exp_m = max(1, int((datetime.datetime(2026, 6, 30) - doj_d).days / 30.4375))
    except Exception:
        exp_m = 36

    prod_emps.append({
        'empNo': emp_no,
        'name': name,
        'dept': dept,
        'section': sec,
        'qualification': 'ITI / Diploma',
        'doj': doj_str,
        'expCount': exp_m,
        'yearExp': f'{int(exp_m/12)}+ years',
        'currentLevel': cur,
        'eligibleLevel': elig
    })

    has_exam = (s_o is not None and isinstance(s_o, (int, float)) and s_o > 0)
    has_gemba = (g_u is not None and isinstance(g_u, (int, float)) and g_u > 0)
    
    quest_tot = int(s_o) if (s_o is not None and isinstance(s_o, (int, float))) else 0
    quest_pct = int(round(s_p * 100)) if (s_p is not None and isinstance(s_p, (int, float))) else 0
    gemba_tot = int(g_u) if (g_u is not None and isinstance(g_u, (int, float))) else 0
    gemba_pct = int(round(g_v * 100)) if (g_v is not None and isinstance(g_v, (int, float))) else 0

    status = 'Passed' if f_x == 'Pass' else ('Passed' if s_q == 'Pass' else 'Failed')

    prod_records[emp_no] = {
        'empNo': emp_no,
        'name': name,
        'dept': dept,
        'section': sec,
        'doj': doj_str,
        'targetLevel': elig,
        'currentLevel': cur,
        'inProgress': False,
        'isCompleted': True if (has_exam or has_gemba) else False,
        'tabSwitchCount': 0,
        'attemptedCount': 40 if elig == 'O' else (30 if elig == 'U' else 20),
        'uMark': quest_tot if elig == 'U' else 0,
        'lMark': quest_tot if elig == 'L' else 0,
        'oMark': quest_tot if elig == 'O' else 0,
        'totalMark': quest_tot,
        'markPct': quest_pct,
        'status': status,
        'attemptDate': '30/06/2026',
        'safetyScore': int(s_j) if isinstance(s_j, (int, float)) else 0,
        'processScore': int(s_k) if isinstance(s_k, (int, float)) else 0,
        'qualityScore': int(s_l) if isinstance(s_l, (int, float)) else 0,
        'ciScore': int(s_m) if isinstance(s_m, (int, float)) else 0,
        'systemScore': int(s_n) if isinstance(s_n, (int, float)) else 0,
        'questTotal': quest_tot,
        'questPct': quest_pct,
        'questStatus': s_q or ('Pass' if quest_pct >= 50 else 'Fail'),
        'gembaSafety': int(g_s) if isinstance(g_s, (int, float)) else 0,
        'gembaSop': int(g_t) if isinstance(g_t, (int, float)) else 0,
        'gembaTotal': gemba_tot,
        'gembaPct': gemba_pct,
        'gembaStatus': g_w or ('Pass' if gemba_pct >= 50 else 'Fail'),
        'finalStatus': f_x or status,
        'proposedLevel': p_y or cur
    }

    if has_gemba:
        prod_ojt[emp_no] = {
            'empNo': emp_no,
            'name': name,
            'section': sec,
            'safetyScore': int(g_s) if isinstance(g_s, (int, float)) else 0,
            'sopScore': int(g_t) if isinstance(g_t, (int, float)) else 0,
            'totalScore': gemba_tot,
            'scorePct': gemba_pct,
            'status': g_w or ('Pass' if gemba_pct >= 50 else 'Fail'),
            'evalDate': '30/06/2026'
        }

print(f'Parsed {len(prod_emps)} Production employees.')
print(f'Generated {len(prod_records)} Production assessment records.')
print(f'Generated {len(prod_ojt)} Production OJT evaluation records.')

# 2. Load existing 283 QA employees from data.js
with open('data.js', encoding='utf-8') as f:
    data_js_content = f.read()

m = re.search(r'const EMPLOYEES = (\[.*?\]);', data_js_content, re.DOTALL)
if not m:
    raise RuntimeError('Could not find const EMPLOYEES in data.js')

qa_employees = json.loads(m.group(1))
# If data.js already had production employees appended previously, filter to QA
qa_only = [e for e in qa_employees if str(e.get('dept', '')).upper() != 'PRODUCTION']
print(f'Original QA employees: {len(qa_only)}')

# Combine QA and Production employees: 283 QA + 1,179 Production = 1,462 Total
all_employees = qa_only + prod_emps
print(f'Combined total employees: {len(all_employees)}')

# 3. Update data.js with all 1,462 employees
employees_json_str = json.dumps(all_employees, indent=2, ensure_ascii=False)
new_data_js = re.sub(
    r'const EMPLOYEES = \[.*?\];',
    f'const EMPLOYEES = {employees_json_str};',
    data_js_content,
    flags=re.DOTALL
)

with open('data.js', 'w', encoding='utf-8') as f:
    f.write(new_data_js)
print('Updated data.js with all 1,462 employees.')

# 4. Update custom_employees.json
with open('custom_employees.json', 'w', encoding='utf-8') as f:
    json.dump(all_employees, f, indent=2, ensure_ascii=False)
print('Updated custom_employees.json with all 1,462 employees.')

# 5. Merge assessment records
records = {}
if os.path.exists('assessment_records_backup_236.json'):
    with open('assessment_records_backup_236.json', encoding='utf-8') as f:
        records.update(json.load(f))
if os.path.exists('assessment_records.json'):
    with open('assessment_records.json', encoding='utf-8') as f:
        records.update(json.load(f))

print(f'Existing QA records: {len(records)}')
records.update(prod_records)
print(f'Merged total assessment records: {len(records)}')

with open('assessment_records.json', 'w', encoding='utf-8') as f:
    json.dump(records, f, indent=2, ensure_ascii=False)
print('Saved assessment_records.json.')

# 6. Merge OJT evaluations
ojt_records = {}
if os.path.exists('ojt_evaluations.json'):
    try:
        with open('ojt_evaluations.json', encoding='utf-8') as f:
            ojt_records = json.load(f)
    except Exception:
        pass

ojt_records.update(prod_ojt)
print(f'Merged total OJT records: {len(ojt_records)}')

with open('ojt_evaluations.json', 'w', encoding='utf-8') as f:
    json.dump(ojt_records, f, indent=2, ensure_ascii=False)
print('Saved ojt_evaluations.json.')

# 7. Update seed_data.js with SEED_RECORDS and SEED_OJT_RECORDS
seed_js_content = f"""// Master Assessment Records - Contains all 1,462 Employees (1,179 Production + 283 QA associates)
(function(global) {{
  'use strict';
  global.SEED_RECORDS = {json.dumps(records, ensure_ascii=False)};
  global.YOKOHAMA_SEED_RECORDS = global.SEED_RECORDS;
  global.SEED_OJT_RECORDS = {json.dumps(ojt_records, ensure_ascii=False)};
  global.YOKOHAMA_SEED_OJT_RECORDS = global.SEED_OJT_RECORDS;
}})(typeof window !== 'undefined' ? window : global);
"""

with open('seed_data.js', 'w', encoding='utf-8') as f:
    f.write(seed_js_content)
print('Saved seed_data.js with full master records.')
