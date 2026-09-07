import openpyxl
import os
import glob
import json
import re
import sys

def get_file_for_section(sec_raw):
    sec = str(sec_raw or '').lower().strip()
    if 'rro' in sec or 'alt' in sec:
        return '83D. ON THE JOB TRAINING EVALUATION -  RRO & ALT Operator.xlsx'
    elif 'preparatory' in sec:
        return '85D. ON THE JOB TRAINING EVALUATION - Preparatory QA.xlsx'
    elif 'solid' in sec:
        return '86D. ON THE JOB TRAINING EVALUATION - Solid Tire QA.xlsx'
    elif 'building' in sec or 'tbm' in sec:
        return '87D. ON THE JOB TRAINING EVALUATION - TBM QA.xlsx'
    elif 'curing' in sec:
        return '88D. ON THE JOB TRAINING EVALUATION - Tire Curing QA.xlsx'
    elif 'warehouse' in sec or 'ware house' in sec or 'data entry' in sec:
        return '89D. ON THE JOB TRAINING EVALUATION -  WAREHOUSE QA.xlsx'
    elif 'finish' in sec:
        return '84D. ON THE JOB TRAINING EVALUATION -  Final Finish Repair Associate.xlsx'
    return '84D. ON THE JOB TRAINING EVALUATION -  Final Finish Repair Associate.xlsx'

def safe_write_cell(ws, row, col, val):
    cell = ws.cell(row=row, column=col)
    if isinstance(cell, openpyxl.cell.cell.MergedCell):
        for rng in ws.merged_cells.ranges:
            if cell.coordinate in rng:
                top_left = rng.coord.split(':')[0]
                ws[top_left].value = val
                return
    else:
        cell.value = val

def update_employee_in_qa_excel(emp_no, ojt_data, qa_folder=r'D:\QA'):
    if not os.path.exists(qa_folder):
        print(f"Directory {qa_folder} does not exist.")
        return False

    sec = ojt_data.get('section', '')
    fname = get_file_for_section(sec)
    fpath = os.path.join(qa_folder, fname)

    if not os.path.exists(fpath):
        print(f"File {fpath} does not exist.")
        return False

    wb = openpyxl.load_workbook(fpath)
    
    # Find employee sheet by emp_no prefix
    target_sheet = None
    emp_prefix = str(emp_no).strip()
    for sname in wb.sheetnames:
        if sname.startswith(emp_prefix):
            target_sheet = wb[sname]
            break

    if target_sheet is None:
        print(f"Sheet for employee {emp_no} not found in {fname}.")
        return False

    ws = target_sheet
    scores = ojt_data.get('scores', {})
    total_score = ojt_data.get('totalScore', 0)
    max_score = ojt_data.get('maxScore', 50)
    pct = ojt_data.get('scorePct', 0)
    comments = ojt_data.get('comments', '')
    safety_rep = ojt_data.get('safetyRep', '')
    quality_rep = ojt_data.get('qualityRep', '')
    ci_rep = ojt_data.get('ciRep', '')
    status = ojt_data.get('qualificationStatus', 'Qualified')

    # Determine row layout
    if fname.startswith('83D'):
        start_cp = 9
        num_cp = 8
        tot_row = 17
        imp_row = 18
        eval_row = 25
        status_row = 34
    elif fname.startswith('84D'):
        start_cp = 9
        num_cp = 7
        tot_row = 16
        imp_row = 17
        eval_row = 24
        status_row = 33
    else:
        start_cp = 8
        num_cp = 10
        tot_row = 18
        imp_row = 19
        eval_row = 26
        status_row = 35

    # Update checkpoint scores
    for i in range(num_cp):
        sno = i + 1
        val = scores.get(str(sno)) or scores.get(sno)
        if val is not None and str(val).isdigit():
            safe_write_cell(ws, start_cp + i, 8, int(val))

    # Total score
    safe_write_cell(ws, tot_row, 8, f"{total_score}/{max_score} = {pct}%")

    # Improvement comments (box starts at column F = 6)
    if comments:
        safe_write_cell(ws, imp_row, 6, comments)

    # Representatives (Safety: col A=1, Quality: col C=3, CI: col I=9)
    if safety_rep:
        safe_write_cell(ws, eval_row, 1, safety_rep)
    if quality_rep:
        safe_write_cell(ws, eval_row, 3, quality_rep)
    if ci_rep:
        safe_write_cell(ws, eval_row, 9, ci_rep)

    # Status (Qualified: col D=4, Not Qualified: col F=6)
    if status == 'Qualified' or (isinstance(pct, (int, float)) and pct >= 70):
        safe_write_cell(ws, status_row, 4, "[X] Qualified")
        safe_write_cell(ws, status_row, 6, "[ ] Not Qualified")
    else:
        safe_write_cell(ws, status_row, 4, "[ ] Qualified")
        safe_write_cell(ws, status_row, 6, "[X] Not Qualified")

    wb.save(fpath)
    print(f"Successfully updated employee {emp_no} in {fname} (Sheet: {ws.title})")
    return True

def sync_all_from_json(json_path=None, qa_folder=r'D:\QA'):
    if json_path is None:
        json_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'ojt_evaluations.json')
    if not os.path.exists(json_path):
        print(f"JSON file {json_path} not found.")
        return
    with open(json_path, 'r', encoding='utf-8') as f:
        all_ojt = json.load(f)
    print(f"Syncing {len(all_ojt)} evaluations to {qa_folder}...")
    for emp_no, data in all_ojt.items():
        update_employee_in_qa_excel(emp_no, data, qa_folder)

if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--emp':
        emp_no = sys.argv[2]
        data = None
        if len(sys.argv) > 3:
            try:
                data = json.loads(sys.argv[3])
            except Exception:
                data = None
        
        if not data:
            json_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'ojt_evaluations.json')
            if os.path.exists(json_path):
                with open(json_path, 'r', encoding='utf-8') as f:
                    all_ojt = json.load(f)
                    data = all_ojt.get(str(emp_no))
        
        if data:
            update_employee_in_qa_excel(emp_no, data)
        else:
            print(f"No data found for employee {emp_no}")
    else:
        sync_all_from_json()
